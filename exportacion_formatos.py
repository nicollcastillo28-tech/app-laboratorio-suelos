"""Diligencia entradas sin volver a guardar ni rediseñar los formatos oficiales.

openpyxl se usa únicamente para leer. La descarga conserva las partes originales
del archivo (incluidos macros, gráficos, estilos y fórmulas) y sustituye solamente
los valores de entrada solicitados por el generador. Las excepciones autorizadas
para sustituir una fórmula se declaran por hoja y celda; actualmente, G21 del 006.
"""
from io import BytesIO
import html
import math
import re
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


class FormatoIncompatibleError(ValueError):
    """La entrada no cabe en el formato original sin alterar su estructura."""


def partes_hojas(contenido):
    with zipfile.ZipFile(BytesIO(contenido)) as z:
        rels = {r.attrib['Id']: r.attrib['Target']
                for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
        hojas = ET.fromstring(z.read('xl/workbook.xml')).find(S + 'sheets')
        resultado = {}
        for hoja in hojas:
            destino = rels.get(hoja.get(R + 'id'))
            if destino:
                resultado[hoja.get('name')] = (destino.lstrip('/') if destino.startswith('/') else 'xl/' + destino)
        return resultado


def escribir_entradas(contenido, parte, valores, *, celdas_formula_editables=()):
    """Cambia entradas; una fórmula sólo se sustituye con autorización explícita por celda."""
    with zipfile.ZipFile(BytesIO(contenido)) as z:
        xml = z.read(parte).decode('utf-8')
        raiz = ET.fromstring(xml)
        combinaciones = [range_boundaries(n.get('ref')) for n in raiz.iter(S + 'mergeCell')]
        existentes = {c.get('r'): c for c in raiz.iter(S + 'c')}
        entradas = {}
        formulas_sustituidas = False
        for ref, valor in valores.items():
            if isinstance(valor, str) and valor.startswith('='):
                raise FormatoIncompatibleError(f'{parte}!{ref}: no se permite introducir una fórmula')
            fila, col = coordinate_to_tuple(ref)
            for c1, f1, c2, f2 in combinaciones:
                if c1 <= col <= c2 and f1 <= fila <= f2 and (col, fila) != (c1, f1):
                    if valor not in (None, ''):
                        raise FormatoIncompatibleError(f'{parte}!{ref}: la casilla está combinada con otra; se conserva el formato original')
                    break
            else:
                original = existentes.get(ref)
                if original is not None and original.find(S + 'f') is not None:
                    if ref not in celdas_formula_editables:
                        if valor not in (None, ''):
                            raise FormatoIncompatibleError(f'{parte}!{ref}: contiene una fórmula del formato original')
                        continue
                    if original.find(S + 'f').get('t') in ('shared', 'array'):
                        raise FormatoIncompatibleError(f'{parte}!{ref}: pertenece a una fórmula compartida o matricial')
                    formulas_sustituidas = True
                if isinstance(valor, float) and not math.isfinite(valor):
                    raise ValueError(f'{parte}!{ref}: valor numérico no finito')
                entradas[ref] = valor

        def reemplazo(ref, atributos, valor):
            # Conservar todos los atributos de formato de la celda, salvo su tipo de dato.
            atributos = re.sub(r'\s+t="[^"]*"', '', atributos)
            atributos = re.sub(r'\s+(?:cm|vm)="[^"]*"', '', atributos)
            if valor is None or valor == '':
                return f'<c r="{ref}"{atributos}/>'
            if isinstance(valor, bool):
                return f'<c r="{ref}"{atributos} t="b"><v>{int(valor)}</v></c>'
            if isinstance(valor, (int, float)):
                return f'<c r="{ref}"{atributos}><v>{valor!r}</v></c>'
            texto = html.escape(str(valor), quote=False)
            return f'<c r="{ref}"{atributos} t="inlineStr"><is><t xml:space="preserve">{texto}</t></is></c>'

        for ref, valor in entradas.items():
            patron = re.compile(r'<c\b\s+r="' + ref + r'"([^>]*?)(?:/>|>.*?</c>)', re.S)
            encontrado = patron.search(xml)
            if encontrado:
                xml = xml[:encontrado.start()] + reemplazo(ref, encontrado.group(1), valor) + xml[encontrado.end():]
                continue
            fila, col = coordinate_to_tuple(ref)
            patron_fila = re.compile(r'<row\b\s+r="' + str(fila) + r'"[^>]*?(?:/>|>(.*?)</row>)', re.S)
            mf = patron_fila.search(xml)
            if mf is None:
                # La ausencia de una fila de entrada es un cambio de formato pendiente de revisar.
                if valor not in (None, ''):
                    raise FormatoIncompatibleError(f'{parte}!{ref}: la fila de entrada no existe en el formato')
                continue
            if mf.group(1) is None:
                if valor in (None, ''):
                    continue
                fila_xml = mf.group(0)[:-2] + '>' + reemplazo(ref, '', valor) + '</row>'
                xml = xml[:mf.start()] + fila_xml + xml[mf.end():]
                continue
            cuerpo = mf.group(1)
            posicion = len(cuerpo)
            for mc in re.finditer(r'<c\b\s+r="([A-Z]+\d+)"', cuerpo):
                if coordinate_to_tuple(mc.group(1))[1] > col:
                    posicion = mc.start()
                    break
            lugar = mf.start(1) + posicion
            xml = xml[:lugar] + reemplazo(ref, '', valor) + xml[lugar:]

        salida = BytesIO()
        with zipfile.ZipFile(salida, 'w', zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                # La cadena de cálculo anterior contiene la fórmula que acabamos de retirar.
                # Excel debe reconstruir esa caché al abrir el libro.
                if formulas_sustituidas and item.filename == 'xl/calcChain.xml':
                    continue
                datos = xml.encode('utf-8') if item.filename == parte else z.read(item.filename)
                if formulas_sustituidas and item.filename == '[Content_Types].xml':
                    datos = re.sub(rb'<Override\b(?=[^>]*\bPartName="/xl/calcChain\.xml")[^>]*/>', b'', datos)
                if formulas_sustituidas and item.filename == 'xl/_rels/workbook.xml.rels':
                    datos = re.sub(rb'<Relationship\b(?=[^>]*\bType="[^"]*/calcChain")[^>]*/>', b'', datos)
                if item.filename == 'xl/workbook.xml':
                    # Invalidar las cachés de cálculo tras llenar entradas, sin cambiar fórmulas.
                    texto = datos.decode('utf-8')
                    def recalcular(m):
                        tag = re.sub(r'\s+(?:fullCalcOnLoad|forceFullCalc)="[^"]*"', '', m.group(0))
                        return tag.replace('/>', ' fullCalcOnLoad="1" forceFullCalc="1"/>')
                    texto = re.sub(r'<calcPr\b[^>]*/>', recalcular, texto)
                    datos = texto.encode('utf-8')
                out.writestr(item, datos)
        return salida.getvalue()


class _HojaEntradas:
    def __init__(self, hoja):
        self._hoja = hoja
        self.entradas = {}

    def __getattr__(self, nombre):
        return getattr(self._hoja, nombre)

    def __getitem__(self, ref):
        return self._hoja[ref]

    def __setitem__(self, ref, valor):
        # Los generadores anteriores reparaban fórmulas del formato anterior.
        # Esas reparaciones quedan desactivadas para conservar el formato nuevo.
        if isinstance(valor, str) and valor.startswith('='):
            if self._hoja[ref].value == valor:
                return
            raise FormatoIncompatibleError(f'{self._hoja.title}!{ref}: se conserva la fórmula original')
        self.entradas[ref] = valor


class LibroFormatoOriginal:
    def __init__(self, path, *, celdas_formula_editables=None, **kwargs):
        with open(path, 'rb') as f:
            self._original = f.read()
        self._libro = load_workbook(BytesIO(self._original), **kwargs)
        self._formulas_editables = celdas_formula_editables or {}
        self._hojas = {h.title: _HojaEntradas(h) for h in self._libro.worksheets}

    @property
    def sheetnames(self):
        return self._libro.sheetnames

    @property
    def worksheets(self):
        return list(self._hojas.values())

    def __getitem__(self, nombre):
        return self._hojas[nombre]

    def __delitem__(self, nombre):
        raise FormatoIncompatibleError(f'No se permite eliminar la hoja original {nombre}')

    def save(self, archivo):
        contenido = self._original
        partes = partes_hojas(contenido)
        for nombre, hoja in self._hojas.items():
            if hoja.entradas:
                contenido = escribir_entradas(contenido, partes[nombre], hoja.entradas,
                                             celdas_formula_editables=self._formulas_editables.get(nombre, ()))
        archivo.write(contenido)
