"""Bitácoras de papel (GDA-FL-xxx) llenadas con lo digitado en la app y entregadas en PDF.

Se parte de la plantilla Excel oficial de cada bitácora (templates/bitacoras): se escriben los valores directamente en el XML de la hoja y se
marcan las casillas (formas dibujadas) editando el XML del dibujo, así se conserva tal cual todo lo demás del formato (logo, formas, líneas,
combinaciones). El PDF sale de convertir ese Excel con LibreOffice, por eso el resultado es igual al formato impreso.
"""
import html
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
import threading
import zipfile
from io import BytesIO

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

DIR_PLANTILLAS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates", "bitacoras")

_LOCK_PDF = threading.Lock()


def to_float(v, default=None):
    try:
        if v is None or str(v).strip() == "":
            return default
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return default


def _num(v, dec=None):
    """Número digitado -> número (para que el Excel lo muestre como número); si no es número, el texto tal cual; vacío -> None."""
    f = to_float(v)
    if f is None:
        t = str(v).strip() if v is not None else ""
        return t or None
    if dec is not None:
        f = round(f, dec)
    return int(f) if f == int(f) and dec is None and abs(f) < 1e15 and "." not in str(v) and "," not in str(v) else f


# ── Lectura de la plantilla (solo para ubicar etiquetas y casillas; nunca se guarda con openpyxl) ─────────────────────────
class Plantilla:
    def __init__(self, archivo):
        self.ruta = os.path.join(DIR_PLANTILLAS, archivo)
        with open(self.ruta, "rb") as f:
            self.bytes = f.read()
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(BytesIO(self.bytes))
        self.ws = wb[wb.sheetnames[0]]
        self.hoja_nombre = wb.sheetnames[0]
        self._merged = {}
        for r in self.ws.merged_cells.ranges:
            for fila in range(r.min_row, r.max_row + 1):
                for col in range(r.min_col, r.max_col + 1):
                    self._merged[(fila, col)] = r
        z = zipfile.ZipFile(BytesIO(self.bytes))
        self.hoja_xml = _ruta_primera_hoja(z)
        self.dibujo_xml = _ruta_dibujo(z, self.hoja_xml)
        self.formas = _leer_formas(z.read(self.dibujo_xml).decode("utf-8")) if self.dibujo_xml else []
        for f in self.formas:   # posición en píxeles (aprox.) para saber qué casilla queda junto a cada etiqueta
            f["x"] = self._x_col(f["col"]) + f["col_off"]
            f["y"] = self._y_fila(f["fila"]) + f["fila_off"]

    def _ancho_col(self, col):
        w = self.ws.column_dimensions[get_column_letter(col)].width
        return (w if w else 8.43) * 7 + 5

    def _alto_fila(self, fila):
        h = self.ws.row_dimensions[fila].height
        return (h if h else 15) * 96 / 72

    def _x_col(self, col):
        return sum(self._ancho_col(c) for c in range(1, col))

    def _y_fila(self, fila):
        return sum(self._alto_fila(r) for r in range(1, fila))

    # Celdas -----------------------------------------------------------------------------------------------------------
    def buscar(self, texto, desde_fila=1, hasta_fila=None, ocurrencia=1):
        """Celda (fila, col) cuyo texto (sin espacios sobrantes, sin mayúsculas) empieza por `texto`."""
        buscado = _norm(texto)
        n = 0
        for fila in self.ws.iter_rows(min_row=desde_fila, max_row=hasta_fila):
            for c in fila:
                if isinstance(c.value, str) and _norm(c.value).startswith(buscado):
                    n += 1
                    if n == ocurrencia:
                        return c.row, c.column
        raise KeyError(f"No hay una celda que diga «{texto}» en la plantilla {os.path.basename(self.ruta)}")

    def valor_de(self, texto, desde_fila=1, hasta_fila=None, ocurrencia=1):
        """Referencia de la celda donde se escribe el dato de la etiqueta `texto`: la primera a su derecha que tenga línea inferior, borde completo
        o esté combinada (las de las bitácoras en papel donde se escribe a mano)."""
        fila, col = self.buscar(texto, desde_fila, hasta_fila, ocurrencia)
        r = self._merged.get((fila, col))
        c = (r.max_col if r else col) + 1
        while c <= col + 30:
            celda = self.ws.cell(row=fila, column=c)
            b = celda.border
            m = self._merged.get((fila, c))
            if celda.value in (None, "") and ((b.bottom is not None and b.bottom.style) or m is not None
                                              or all(getattr(b, s) is not None and getattr(b, s).style for s in ("left", "right", "top", "bottom"))):
                return f"{get_column_letter(m.min_col if m else c)}{m.min_row if m else fila}"
            if celda.value not in (None, ""):
                break
            c += 1
        # Sin línea de escribir en la plantilla: se escribe después de donde termina el texto de la etiqueta (que se desborda a la derecha)
        ancho_etq, c = 0, col
        necesario = len(str(self.ws.cell(row=fila, column=col).value or "")) * 2.2
        while c < col + 30:
            ancho_etq += self._ancho_col(c) / 7
            c += 1
            if ancho_etq >= necesario:
                break
        celda = self.ws.cell(row=fila, column=c)
        if celda.value in (None, ""):
            m = self._merged.get((fila, c))
            return f"{get_column_letter(m.min_col if m else c)}{m.min_row if m else fila}"
        raise KeyError(f"No se halló dónde escribir «{texto}»")

    def _subrayada(self, fila, col):
        c = self.ws.cell(row=fila, column=col)
        b = c.border
        return (self._merged.get((fila, col)) is not None) or (b.bottom is not None and b.bottom.style) is not None and bool(b.bottom and b.bottom.style)

    def _fin_linea(self, fila, col):
        """Última columna de la línea de escribir que empieza en `col`: la última celda con línea inferior antes de la próxima celda con texto."""
        tope = self.ws.max_column
        pa = self.ws.print_area
        if pa:
            m = re.search(r"\$([A-Z]+)\$\d+$", pa if isinstance(pa, str) else str(pa))
            if m:
                tope = column_index_from_string(m.group(1))
        ultimo, c = col, col
        while c <= min(tope, col + 40):
            v = self.ws.cell(row=fila, column=c).value
            if c > col and isinstance(v, str) and v.strip():
                break
            if c > col and v not in (None, "") and not isinstance(v, str):
                break
            if self._subrayada(fila, c):
                ultimo = c
            c += 1
        return ultimo

    def lineas(self, ref, texto, max_lineas=4, factor=1.15):
        """Reparte un texto largo (descripción, observaciones) en las líneas de escribir que hay desde `ref` hacia abajo, cortando por palabras
        según lo ancho de la línea. Si no cabe, la última línea se combina y el texto se encoge para que se lea completo."""
        import textwrap
        m = re.match(r"([A-Z]+)(\d+)$", ref)
        col, fila = column_index_from_string(m.group(1)), int(m.group(2))
        r = self._merged.get((fila, col))
        if r:
            ancho, c_fin = sum(self._ancho_col(c) for c in range(r.min_col, r.max_col + 1)) / 7, r.max_col
        else:
            c_fin = self._fin_linea(fila, col)
            ancho = sum(self._ancho_col(c) for c in range(col, c_fin + 1)) / 7
        por_linea = max(12, int(ancho * factor))
        filas = [fila]
        f = fila + 1
        while len(filas) < max_lineas and self._subrayada(f, col) and all(
                self.ws.cell(row=f, column=cc).value in (None, "") or not str(self.ws.cell(row=f, column=cc).value).strip() for cc in range(1, col)):
            filas.append(f)
            f += 1
        trozos = textwrap.wrap(" ".join(str(texto).split()), por_linea) or [""]
        salida = {}
        for k, fila_k in enumerate(filas):
            if k < len(filas) - 1 and k < len(trozos):
                salida[f"{get_column_letter(col)}{fila_k}"] = Libre(trozos[k])
            elif k == len(filas) - 1 and k < len(trozos):
                resto = " ".join(trozos[k:])
                ref_k = f"{get_column_letter(col)}{fila_k}"
                if len(trozos) - k > 1 and not r:
                    v = Ajustado(resto)
                    v.fusion = f"{ref_k}:{get_column_letter(c_fin)}{fila_k}" if c_fin > col else None
                    salida[ref_k] = v
                else:
                    salida[ref_k] = Libre(resto)
        return salida

    def celda_a_la_derecha(self, ref, saltos=1):
        m = re.match(r"([A-Z]+)(\d+)$", ref)
        col, fila = column_index_from_string(m.group(1)), int(m.group(2))
        r = self._merged.get((fila, col))
        c = (r.max_col if r else col) + saltos
        mm = self._merged.get((fila, c))
        return f"{get_column_letter(mm.min_col if mm else c)}{mm.min_row if mm else fila}"

    # Casillas ---------------------------------------------------------------------------------------------------------
    def casillas_en_celda(self, fila, col):
        """Ids de las casillas dibujadas sobre una celda (o su rango combinado), de izquierda a derecha."""
        r = self._merged.get((fila, col))
        f_ini, f_fin = (r.min_row, r.max_row) if r else (fila, fila)
        c_ini, c_fin = (r.min_col, r.max_col) if r else (col, col)
        x_ini, x_fin = self._x_col(c_ini), self._x_col(c_fin + 1)
        y_ini, y_fin = self._y_fila(f_ini), self._y_fila(f_fin + 1)
        dentro = [f for f in self.formas if y_ini - 6 <= f["y"] + f["alto"] / 2 <= y_fin + 6 and x_ini - 4 <= f["x"] <= x_fin + 40]
        return [f["id"] for f in sorted(dentro, key=lambda f: f["x"])]

    def casilla_de(self, texto_etiqueta, ocurrencia=1, desde_fila=1, hasta_fila=None):
        """Id de la forma (casilla) que acompaña a la etiqueta: la que está en su misma línea y más cerca de su borde derecho (las casillas de las
        bitácoras van justo después del texto)."""
        fila, col = self.buscar(texto_etiqueta, desde_fila, hasta_fila, ocurrencia)
        r = self._merged.get((fila, col))
        f_ini, f_fin = (r.min_row, r.max_row) if r else (fila, fila)
        c_ini, c_fin = (r.min_col, r.max_col) if r else (col, col)
        x_ini, x_fin = self._x_col(c_ini), self._x_col(c_fin + 1)
        y_ini, y_fin = self._y_fila(f_ini), self._y_fila(f_fin + 1)
        mejor = None
        for f in self.formas:
            cy = f["y"] + f["alto"] / 2
            if not (y_ini - 6 <= cy <= y_fin + 6) or f["x"] < x_ini - 4:
                continue
            dist = abs(f["x"] - x_fin)
            if mejor is None or dist < mejor[0]:
                mejor = (dist, f["id"])
        if mejor is None:
            raise KeyError(f"No hay casilla junto a «{texto_etiqueta}»")
        return mejor[1]


def _norm(t):
    t = unicodedata.normalize("NFD", str(t))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", t).strip().upper()


def _ruta_primera_hoja(z):
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid = re.search(r'<sheet [^>]*r:id="([^"]+)"', wb).group(1)
    for m in re.finditer(r"<Relationship\b[^>]*>", rels):
        tag = m.group(0)
        if f'Id="{rid}"' in tag:
            destino = re.search(r'Target="([^"]+)"', tag).group(1)
            return "xl/" + destino.lstrip("/").replace("xl/", "", 1) if not destino.startswith("/xl/") else destino.lstrip("/")
    raise ValueError("No se encontró la primera hoja")


def _ruta_dibujo(z, hoja_xml):
    rels = hoja_xml.replace("worksheets/", "worksheets/_rels/") + ".rels"
    if rels not in z.namelist():
        return None
    for m in re.finditer(r"<Relationship\b[^>]*>", z.read(rels).decode("utf-8")):
        if "/drawing" in m.group(0) and "vmlDrawing" not in m.group(0):
            destino = re.search(r'Target="([^"]+)"', m.group(0)).group(1)
            return "xl/" + destino.replace("../", "") if not destino.startswith("/") else destino.lstrip("/")
    return None


def _leer_formas(xml):
    """Casillas del dibujo de la hoja: formas (cuadros) o imágenes pequeñas de un cuadro vacío."""
    formas = []
    for m in re.finditer(r"<xdr:(oneCellAnchor|twoCellAnchor)\b[^>]*>(.*?)</xdr:\1>", xml, re.S):
        cuerpo = m.group(2)
        es_forma, es_imagen = ("<xdr:sp " in cuerpo or "<xdr:sp>" in cuerpo), "<xdr:pic" in cuerpo
        if not (es_forma or es_imagen):
            continue
        desde = re.search(r"<xdr:from>.*?</xdr:from>", cuerpo, re.S)
        col0 = int(re.search(r"<xdr:col>(\d+)</xdr:col>", desde.group(0)).group(1))
        col_off = int(re.search(r"<xdr:colOff>(-?\d+)</xdr:colOff>", desde.group(0)).group(1))
        fila0 = int(re.search(r"<xdr:row>(\d+)</xdr:row>", desde.group(0)).group(1))
        fila_off = int(re.search(r"<xdr:rowOff>(-?\d+)</xdr:rowOff>", desde.group(0)).group(1))
        cid = int(re.search(r'<xdr:cNvPr id="(\d+)"', cuerpo).group(1))
        ext = re.search(r'<xdr:ext cx="(\d+)" cy="(\d+)"', cuerpo) or re.search(r'<a:ext cx="(\d+)" cy="(\d+)"', cuerpo)
        ancho, alto = (int(ext.group(1)), int(ext.group(2))) if ext else (0, 0)
        if ancho > 1_200_000 or alto > 1_200_000 or ancho < 15 * 9525 or alto < 15 * 9525:   # logos, diagramas, rayitas
            continue
        formas.append({"id": cid, "col": col0 + 1, "fila": fila0 + 1, "col_off": col_off // 9525, "fila_off": fila_off // 9525,
                       "ancho": ancho // 9525, "alto": alto // 9525, "imagen": es_imagen and not es_forma,
                       "desde_xml": desde.group(0), "ext_emu": (ancho, alto)})
    return formas


# ── Escritura en el XML ──────────────────────────────────────────────────────────────────────────────────────────────────
def _col_num(letras):
    return column_index_from_string(letras)


class Izq(str):
    """Texto que se alinea a la izquierda dentro de su casilla (con ajuste de línea)."""


class Libre(str):
    """Texto a la izquierda, sin ajustar ni encoger: se desborda sobre la línea de escribir de la bitácora (como se escribiría a mano)."""


class Ajustado(Libre):
    """Texto largo que no cabe en las líneas disponibles: se combina la línea (`fusion`) y se encoge para que se lea completo."""
    fusion = None


def _celda_xml(ref, estilo, valor):
    attr_s = f' s="{estilo}"' if estilo else ""
    if valor is None or valor == "":
        return f'<c r="{ref}"{attr_s}/>'
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        texto = html.escape(str(valor), quote=False)
        return f'<c r="{ref}"{attr_s} t="inlineStr"><is><t xml:space="preserve">{texto}</t></is></c>'
    return f'<c r="{ref}"{attr_s}><v>{repr(float(valor)) if isinstance(valor, float) else valor}</v></c>'


class _Estilos:
    """Estilos nuevos (copias del de la celda con otra alineación) agregados a styles.xml."""

    def __init__(self, xml):
        self.xml = xml
        m = re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", xml, re.S)
        self.xfs = re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", m.group(1), re.S) if m else []
        self.nuevos = {}

    def alineado(self, indice, modo, numerico=False):
        clave = (indice, modo, numerico)
        if clave in self.nuevos:
            return self.nuevos[clave]
        i = int(indice or 0)
        base = self.xfs[i] if i < len(self.xfs) else '<xf numFmtId="0" fontId="0" fillId="0" borderId="0"/>'
        base = re.sub(r'\s*applyAlignment="\d"', "", base)
        if numerico:   # los números se muestran tal cual (algunas casillas de la plantilla traen formato de fracción o de fecha)
            base = re.sub(r'\s*(numFmtId|applyNumberFormat)="\d+"', "", base)
            base = base.replace("<xf ", '<xf numFmtId="0" applyNumberFormat="1" ', 1)
        izq = modo in ("izq", "libre", "ajustado")
        al = (f'<alignment horizontal="{"left" if izq else "center"}" vertical="center" wrapText="{1 if modo == "izq" else 0}" '
              f'shrinkToFit="{0 if modo in ("izq", "libre") else 1}"/>')
        if base.endswith("/>"):
            xf = base[:-2] + ' applyAlignment="1">' + al + "</xf>"
        else:
            xf = re.sub(r"<alignment\b[^>]*/>", "", base)
            xf = re.sub(r"^(<xf\b[^>]*?)>", r'\1 applyAlignment="1">', xf, count=1)
            xf = xf.replace("</xf>", al + "</xf>")
        self.xfs.append(xf)
        self.nuevos[clave] = len(self.xfs) - 1
        return self.nuevos[clave]

    def aplicar(self):
        if not self.nuevos:
            return self.xml
        cuerpo = "".join(self.xfs)
        return re.sub(r"<cellXfs\b[^>]*>.*?</cellXfs>", lambda m: f'<cellXfs count="{len(self.xfs)}">' + cuerpo + "</cellXfs>", self.xml,
                      count=1, flags=re.S)


def escribir_celdas(xml, celdas, estilos=None):
    """Escribe `celdas` ({"D6": valor}) en el XML de una hoja conservando el estilo de cada celda, con alineación centrada (o a la izquierda
    si el valor es un `Izq`)."""
    for ref, valor in celdas.items():
        if valor is None:
            continue
        m_ref = re.match(r"([A-Z]+)(\d+)$", ref)
        col, fila = m_ref.group(1), int(m_ref.group(2))
        patron = re.compile(r'<c r="' + ref + r'"([^>]*?)(?:/>|>.*?</c>)', re.S)
        m = patron.search(xml)
        if m:
            est = re.search(r'\bs="(\d+)"', m.group(1))
            est = est.group(1) if est else ""
            if estilos is not None:
                est = str(estilos.alineado(est, "ajustado" if isinstance(valor, Ajustado) else "libre" if isinstance(valor, Libre)
                                       else "izq" if isinstance(valor, Izq) else "centro",
                                       numerico=isinstance(valor, (int, float)) and not isinstance(valor, bool)))
            xml = xml[:m.start()] + _celda_xml(ref, est, valor) + xml[m.end():]
            continue
        m_fila = re.search(r'<row r="' + str(fila) + r'"[^>]*?(?:/>|>(.*?)</row>)', xml, re.S)
        if not m_fila or m_fila.group(1) is None:
            continue
        cuerpo, inicio = m_fila.group(1), m_fila.start(1)
        pos = len(cuerpo)
        for mc in re.finditer(r'<c r="([A-Z]+)\d+"', cuerpo):
            if _col_num(mc.group(1)) > _col_num(col):
                pos = mc.start()
                break
        xml = xml[:inicio + pos] + _celda_xml(ref, "", valor) + xml[inicio + pos:]
    return xml


_MARCA = ('<a:p><a:pPr marL="0" lvl="0" indent="0" algn="ctr" rtl="0"><a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/>'
          '</a:spcAft><a:buNone/></a:pPr><a:r><a:rPr lang="es-CO" sz="__SZ__" b="1"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:rPr>'
          '<a:t>X</a:t></a:r></a:p>')


def marcar_formas(xml, ids, formas=()):
    """Pone una X (centrada) dentro de las formas con esos ids; si la casilla es una imagen, se superpone un cuadro de texto con la X."""
    por_id = {f["id"]: f for f in formas}
    siguiente = max([int(x) for x in re.findall(r'<xdr:cNvPr id="(\d+)"', xml)] + [0]) + 1
    extra = ""
    for cid in ids:
        f = por_id.get(cid)
        if f and f.get("imagen"):
            cx, cy = f["ext_emu"]
            sz = max(1000, min(4000, int(cy / 12700 * 0.62 * 100)))
            extra += (f'<xdr:oneCellAnchor>{f["desde_xml"]}<xdr:ext cx="{cx}" cy="{cy}"/><xdr:sp macro="" textlink=""><xdr:nvSpPr>'
                      f'<xdr:cNvPr id="{siguiente}" name="Marca {siguiente}"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/>'
                      f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></xdr:spPr>'
                      f'<xdr:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/><a:lstStyle/>'
                      + _MARCA.replace("__SZ__", str(sz)) + '</xdr:txBody></xdr:sp><xdr:clientData/></xdr:oneCellAnchor>')
            siguiente += 1
            continue
        patron = re.compile(r'(<xdr:sp\b(?:(?!</xdr:sp>).)*?<xdr:cNvPr id="' + str(cid) + r'"(?:(?!</xdr:sp>).)*?)(</xdr:sp>)', re.S)
        m = patron.search(xml)
        if not m:
            continue
        cuerpo = m.group(1)
        cuerpo = re.sub(r"<a:p>.*?</a:p>", "", cuerpo, flags=re.S) if "<a:p>" in cuerpo else cuerpo
        alto = re.search(r'<a:ext cx="\d+" cy="(\d+)"', cuerpo)
        sz = max(1000, min(4000, int(int(alto.group(1)) / 12700 * 0.62 * 100))) if alto else 1400   # la X ocupa ~60 % de la casilla
        if "</xdr:txBody>" in cuerpo:
            cuerpo = cuerpo.replace("</xdr:txBody>", _MARCA.replace("__SZ__", str(sz)) + "</xdr:txBody>")
        cuerpo = re.sub(r'anchor="t"', 'anchor="ctr"', cuerpo)
        cuerpo = re.sub(r'(lIns|rIns)="\d+"', r'\1="0"', cuerpo)
        xml = xml[:m.start()] + cuerpo + m.group(2) + xml[m.end():]
    if extra:
        xml = xml.replace("</xdr:wsDr>", extra + "</xdr:wsDr>")
    return xml


def _se_cruza(plantilla, rango):
    a, b = rango.split(":")
    ma, mb = re.match(r"([A-Z]+)(\d+)", a), re.match(r"([A-Z]+)(\d+)", b)
    fila = int(ma.group(2))
    return any(plantilla._merged.get((fila, c)) for c in range(column_index_from_string(ma.group(1)), column_index_from_string(mb.group(1)) + 1))


def llenar(plantilla, celdas, casillas=()):
    """Excel (bytes) de la plantilla con las celdas escritas y las casillas marcadas."""
    z_in = zipfile.ZipFile(BytesIO(plantilla.bytes))
    estilos = _Estilos(z_in.read("xl/styles.xml").decode("utf-8"))
    celdas = dict(celdas)
    for ref, valor in list(celdas.items()):   # un texto libre se desborda a la derecha solo si las celdas vecinas están realmente vacías
        if isinstance(valor, Libre):
            m = re.match(r"([A-Z]+)(\d+)$", ref)
            col, fila = column_index_from_string(m.group(1)), int(m.group(2))
            for c in range(col + 1, col + 40):
                v = plantilla.ws.cell(row=fila, column=c).value
                if isinstance(v, str) and not v.strip() and f"{get_column_letter(c)}{fila}" not in celdas:
                    celdas[f"{get_column_letter(c)}{fila}"] = ""
                elif v not in (None, ""):
                    break
    hoja = escribir_celdas(z_in.read(plantilla.hoja_xml).decode("utf-8"), celdas, estilos)
    for valor in celdas.values():
        rango = getattr(valor, "fusion", None) if isinstance(valor, Ajustado) else None
        if rango and hoja.count("<mergeCells") and not _se_cruza(plantilla, rango):
            hoja = re.sub(r'<mergeCells count="(\d+)">', lambda m: f'<mergeCells count="{int(m.group(1)) + 1}">', hoja, count=1)
            hoja = hoja.replace("</mergeCells>", f'<mergeCell ref="{rango}"/></mergeCells>', 1)
    bio = BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z_out:
        for item in z_in.infolist():
            datos = z_in.read(item.filename)
            if item.filename == plantilla.hoja_xml:
                datos = hoja.encode("utf-8")
            elif item.filename == "xl/styles.xml":
                datos = estilos.aplicar().encode("utf-8")
            elif plantilla.dibujo_xml and item.filename == plantilla.dibujo_xml and casillas:
                datos = marcar_formas(datos.decode("utf-8"), casillas, plantilla.formas).encode("utf-8")
            z_out.writestr(item, datos)
    return bio.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────────────────────────────────────────────────
def _buscar_soffice():
    for nombre in ("soffice", "libreoffice"):
        ruta = shutil.which(nombre)
        if ruta:
            return ruta
    for ruta in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                 "/usr/bin/soffice", "/usr/lib/libreoffice/program/soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(ruta):
            return ruta
    return None


def libreoffice_disponible():
    return _buscar_soffice() is not None


def _dir_temporal():
    base = os.environ.get("GEODELTA_TMP")
    if base:
        os.makedirs(base, exist_ok=True)
    return tempfile.mkdtemp(dir=base or None, prefix="bit_")


def xlsx_a_pdf(xlsx_bytes, paginas="1"):
    """Convierte el Excel a PDF con LibreOffice (solo las páginas indicadas: la primera es el formato, las demás son hojas auxiliares)."""
    soffice = _buscar_soffice()
    if not soffice:
        raise RuntimeError("No está instalado LibreOffice en el servidor (ver packages.txt).")
    with _LOCK_PDF:
        tmp = _dir_temporal()
        try:
            origen = os.path.join(tmp, "bitacora.xlsx")
            with open(origen, "wb") as f:
                f.write(xlsx_bytes)
            perfil = os.path.join(os.environ.get("GEODELTA_TMP") or tempfile.gettempdir(), "geodelta_lo_perfil")
            filtro = 'pdf:calc_pdf_Export:{"PageRange":{"type":"string","value":"' + paginas + '"}}'
            cmd = [soffice, f"-env:UserInstallation={_a_uri(perfil)}", "--headless", "--norestore", "--convert-to", filtro, "--outdir", tmp, origen]
            res = subprocess.run(cmd, capture_output=True, timeout=180)
            salida = os.path.join(tmp, "bitacora.pdf")
            if not os.path.exists(salida):
                raise RuntimeError("LibreOffice no generó el PDF: " + (res.stderr or res.stdout or b"").decode("utf-8", "ignore")[-300:])
            with open(salida, "rb") as f:
                return f.read()
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


def _a_uri(ruta):
    ruta = os.path.abspath(ruta).replace("\\", "/")
    return "file:///" + ruta.lstrip("/") if re.match(r"^[A-Za-z]:", ruta) else "file://" + ruta


def unir_pdfs(lista):
    from pypdf import PdfWriter, PdfReader
    w = PdfWriter()
    for b in lista:
        for p in PdfReader(BytesIO(b)).pages:
            w.add_page(p)
    out = BytesIO()
    w.write(out)
    return out.getvalue()
