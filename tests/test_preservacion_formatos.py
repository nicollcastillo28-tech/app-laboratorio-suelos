"""Pruebas de conservación de los formatos: python -m unittest discover -s tests."""
from pathlib import Path
from io import BytesIO
import hashlib
import json
import re
import unittest
import ast
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from exportacion_formatos import escribir_entradas, partes_hojas, FormatoIncompatibleError, LibroFormatoOriginal

ROOT = Path(__file__).resolve().parents[1]
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


class FormatosOficialesTest(unittest.TestCase):
    def test_archivos_identicos_a_los_originales_entregados(self):
        registros = json.loads((ROOT / 'templates/formatos_origen.json').read_text(encoding='utf-8'))
        self.assertEqual(len(registros), 19)
        for registro in registros:
            with self.subTest(formato=registro['destino']):
                digest = hashlib.sha256((ROOT / 'templates' / registro['destino']).read_bytes()).hexdigest()
                self.assertEqual(digest, registro['sha256'])

    def test_llenar_humedad_sin_alterar_formulas_ni_diseno(self):
        original = (ROOT / 'templates/GDA-FLC-014_humedad_natural.xlsx').read_bytes()
        parte = partes_hojas(original)['Hoja1']
        valores = {'D6': 'CLIENTE DE PRUEBA', 'I19': 'R-917', 'I20': 137.28, 'I21': 113.56, 'I22': 19.84}
        resultado = escribir_entradas(original, parte, valores)
        libro = load_workbook(BytesIO(resultado))
        for celda, valor in valores.items():
            self.assertEqual(libro['Hoja1'][celda].value, valor)
        with zipfile.ZipFile(BytesIO(original)) as a, zipfile.ZipFile(BytesIO(resultado)) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for nombre in a.namelist():
                if nombre not in (parte, 'xl/workbook.xml', 'xl/styles.xml'):
                    self.assertEqual(a.read(nombre), b.read(nombre), nombre)
            antes, despues = a.read(parte), b.read(parte)
            sin_datos = lambda xml: re.sub(rb'<conditionalFormatting\b.*?</conditionalFormatting>', b'', re.sub(rb'<sheetData>.*?</sheetData>', b'<sheetData/>', xml, flags=re.S), flags=re.S)
            self.assertEqual(sin_datos(antes), sin_datos(despues))
            def formulas(xml):
                return {c.get('r'): ET.tostring(c.find(S+'f')) for c in ET.fromstring(xml).iter(S+'c') if c.find(S+'f') is not None}
            self.assertEqual(formulas(antes), formulas(despues))

    def test_cero_es_numero_y_vacio_es_celda_vacia(self):
        original = (ROOT / 'templates/GDA-FLC-014_humedad_natural.xlsx').read_bytes()
        parte = partes_hojas(original)['Hoja1']
        resultado = escribir_entradas(original, parte, {'I19': '', 'I20': 0, 'I21': None})
        ws = load_workbook(BytesIO(resultado))['Hoja1']
        self.assertIsNone(ws['I19'].value)
        self.assertEqual(ws['I20'].value, 0)
        self.assertEqual(ws['I20'].data_type, 'n')
        self.assertIsNone(ws['I21'].value)

    def test_no_sobrescribir_formula_original(self):
        original = (ROOT / 'templates/GDA-FLC-014_humedad_natural.xlsx').read_bytes()
        with self.assertRaises(FormatoIncompatibleError):
            escribir_entradas(original, partes_hojas(original)['Hoja1'], {'I24': 25.3})

    def test_no_descombinar_casillas_cbr(self):
        original = (ROOT / 'templates/GDA-FLC-013_cbr.xlsx').read_bytes()
        with self.assertRaises(FormatoIncompatibleError):
            escribir_entradas(original, partes_hojas(original)['FORMATO'], {'I39': 4540})

    def test_no_admitir_formulas_como_datos(self):
        original = (ROOT / 'templates/GDA-FLC-014_humedad_natural.xlsx').read_bytes()
        with self.assertRaises(FormatoIncompatibleError):
            escribir_entradas(original, partes_hojas(original)['Hoja1'], {'I20': '=1+2'})

    def test_gravedad_g21_recibe_dato_app_y_conserva_resto_del_formato(self):
        # Ejecutar los dos generadores reales, sin iniciar Streamlit ni consultar Supabase.
        ruta = ROOT / 'templates/GDA-FLC-006_gravedad_arcilla.xlsx'
        original = ruta.read_bytes()
        arbol = ast.parse((ROOT / 'app.py').read_text(encoding='utf-8'))
        nombres = {'_generar_excel_gravedad', 'generar_excel_gravedad_arcilla'}
        funciones = [n for n in arbol.body if isinstance(n, ast.FunctionDef) and n.name in nombres]
        for funcion in funciones:
            funcion.decorator_list = []
        ns = {'BytesIO': BytesIO, 'TEMPLATE_GESP_ARCILLA': str(ruta),
              '_cargar_plantilla_informe': LibroFormatoOriginal,
              'to_float': lambda v: None if v in (None, '') else float(v),
              '_fecha_ddmmaaaa': lambda v: v, '_fecha_recepcion_perf': lambda *a: '12/10/2026',
              'fecha_ejecucion_informe': lambda *a: '13/10/2026', 'fecha_emision_hoy': lambda: '16/10/2026',
              'get_perforacion': lambda *a: {'tipo': 'Sondeo'}, 'TIPO_PERFORACION_EXCEL': {'Sondeo': 'SONDEO'},
              'descripcion_para_informe': lambda m: 'SUELO DE PRUEBA'}
        exec(compile(ast.Module(body=funciones, type_ignores=[]), str(ROOT/'app.py'), 'exec'), ns)
        muestra = {'numero': '13', 'profundidad_de': '2.31', 'profundidad_hasta': '3.47'}
        def formulas(xml):
            return {c.get('r'): ET.tostring(c.find(S+'f')) for c in ET.fromstring(xml).iter(S+'c') if c.find(S+'f') is not None}
        for digitado, esperado in [('691.23', 691.23), ('0', 0), ('', None)]:
            with self.subTest(digitado=digitado):
                datos = {'gesp_a_picnometro': '2', 'gesp_a_masa_aire': '92.34',
                         'gesp_a_masa_pic_agua': digitado, 'gesp_a_masa_pic_muestra': '747.81',
                         'gesp_a_temp': '22.5', 'gesp_a_pct_retenido': '0'}
                resultado = ns['generar_excel_gravedad_arcilla']('GDA-998-26', 'S7', muestra,
                                                                {'nombre': 'PRUEBA'}, datos)
                ws = load_workbook(BytesIO(resultado))['FORMATO']
                self.assertEqual(ws['G21'].value, esperado)
                self.assertNotEqual(ws['G21'].data_type, 'f')
                for celda, valor in {'G19': 2, 'G20': 92.34, 'G22': 747.81, 'G23': 22.5, 'D32': 0}.items():
                    self.assertEqual(ws[celda].value, valor)
                with zipfile.ZipFile(BytesIO(original)) as a, zipfile.ZipFile(BytesIO(resultado)) as b:
                    parte = partes_hojas(original)['FORMATO']
                    antes, despues = a.read(parte), b.read(parte)
                    fa, fb = formulas(antes), formulas(despues)
                    self.assertIn('G21', fa)
                    del fa['G21']
                    self.assertEqual(fa, fb)
                    sin_datos = lambda xml: re.sub(rb'<conditionalFormatting\b.*?</conditionalFormatting>', b'', re.sub(rb'<sheetData>.*?</sheetData>', b'<sheetData/>', xml, flags=re.S), flags=re.S)
                    self.assertEqual(sin_datos(antes), sin_datos(despues))
                    origen_ws = load_workbook(BytesIO(original))['FORMATO']
                    self.assertEqual(origen_ws['G21']._style, ws['G21']._style)
                    excepciones = {parte, 'xl/styles.xml', 'xl/workbook.xml', 'xl/calcChain.xml', '[Content_Types].xml', 'xl/_rels/workbook.xml.rels'}
                    for nombre in a.namelist():
                        if nombre not in excepciones:
                            self.assertEqual(a.read(nombre), b.read(nombre), nombre)
                    self.assertNotIn('xl/calcChain.xml', b.namelist())
                    self.assertNotIn(b'/calcChain', b.read('xl/_rels/workbook.xml.rels'))
                    self.assertNotIn(b'/xl/calcChain.xml', b.read('[Content_Types].xml'))
                self.assertEqual(ruta.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
