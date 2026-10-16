from io import BytesIO
from pathlib import Path
import re, unittest, zipfile, xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import range_boundaries,get_column_letter
from exportacion_formatos import escribir_entradas,partes_hojas
from presentacion_excel import ajustar_presentacion,reparar_formulas_clasificacion

ROOT=Path(__file__).resolve().parents[1]
S='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'

def celdas(rango):
    c1,r1,c2,r2=range_boundaries(rango)
    return [f'{get_column_letter(c)}{r}' for r in range(r1,r2+1) for c in range(c1,c2+1)]


class PresentacionTest(unittest.TestCase):
    def test_humedad_reglas_consultan_datos_correctos_y_sus_combinaciones(self):
        original=(ROOT/'templates/GDA-FLC-014_humedad_natural.xlsx').read_bytes()
        parte=partes_hojas(original)['Hoja1']
        valores={'D6':'CLIENTE','D12':'SONDEO','D13':'ARCILLA','I19':'R-17','I20':0,'I21':37.15,'I22':17.12}
        salida=escribir_entradas(original,parte,valores)
        ws=load_workbook(BytesIO(salida))['Hoja1']
        with zipfile.ZipFile(BytesIO(original)) as a,zipfile.ZipFile(BytesIO(salida)) as b:
            raiz=ET.fromstring(b.read(parte));antes=ET.fromstring(a.read(parte))
            cobertura=lambda root:{celda for cf in root.iter(S+'conditionalFormatting') for rango in cf.get('sqref').split() for celda in celdas(rango)}
            self.assertEqual(cobertura(antes),cobertura(raiz))
            for cf in raiz.iter(S+'conditionalFormatting'):
                origen=cf.get('sqref').split()[0].split(':')[0]
                for regla in cf:
                    formula=regla.findtext(S+'formula') or ''
                    if not formula.startswith('LEN(TRIM('):
                        continue
                    self.assertEqual(regla.get('type'),'expression')
                    for rango in cf.get('sqref').split():
                        for celda in celdas(rango):
                            trasladada=Translator('='+formula,origin=origen).translate_formula(celda)
                            destino=re.fullmatch(r'=LEN\(TRIM\((\$?[A-Z]+\$?\d+)\)\)=0',trasladada).group(1).replace('$','')
                            esperado=next((str(r).split(':')[0] for r in ws.merged_cells.ranges if celda in r),celda)
                            self.assertEqual(destino,esperado,(celda,formula,destino,esperado))
                            if destino in valores:
                                self.assertFalse(str(ws[destino].value or (0 if ws[destino].value==0 else '')).strip()=='')
            regla_alerta=[r for r in raiz.iter(S+'cfRule') if r.get('type')=='cellIs']
            self.assertEqual(len(regla_alerta),1)
            self.assertEqual(regla_alerta[0].get('operator'),'greaterThan')
            self.assertEqual(regla_alerta[0].findtext(S+'formula'),'$S$31')
            self.assertEqual(a.read('xl/styles.xml').split(b'<cellXfs')[0],b.read('xl/styles.xml').split(b'<cellXfs')[0])
        for celda in valores:
            self.assertEqual(ws[celda].alignment.horizontal,'center')
            self.assertEqual(ws[celda].alignment.vertical,'center')
        self.assertEqual((ROOT/'templates/GDA-FLC-014_humedad_natural.xlsx').read_bytes(),original)

    def test_centrado_conserva_formato_numerico_bordes_fuente_y_ajuste_texto(self):
        original=(ROOT/'templates/GDA-FLC-014_humedad_natural.xlsx').read_bytes()
        parte=partes_hojas(original)['Hoja1'];refs=['D12','D13','I20']
        salida=ajustar_presentacion(original,parte,refs)
        antes=load_workbook(BytesIO(original))['Hoja1'];despues=load_workbook(BytesIO(salida))['Hoja1']
        for ref in refs:
            a,b=antes[ref],despues[ref]
            self.assertEqual(a.number_format,b.number_format)
            self.assertEqual(a.font.copy(),b.font.copy());self.assertEqual(a.border.copy(),b.border.copy())
            self.assertEqual(a.fill.copy(),b.fill.copy());self.assertEqual(a.protection.copy(),b.protection.copy())
            self.assertEqual(a.alignment.wrapText,b.alignment.wrapText)
        with zipfile.ZipFile(BytesIO(salida)) as a,zipfile.ZipFile(BytesIO(ajustar_presentacion(salida,parte,refs))) as b:
            self.assertEqual(a.read('xl/styles.xml'),b.read('xl/styles.xml'))

    def test_granulometria_formulas_autorizadas_sin_cambiar_objetos_ni_aspecto(self):
        original=(ROOT/'templates/CLASIFICACION_DE_SUELOS.xlsm').read_bytes()
        parte=partes_hojas(original)['FORMATO'];salida=reparar_formulas_clasificacion(original,parte)
        with zipfile.ZipFile(BytesIO(original)) as a,zipfile.ZipFile(BytesIO(salida)) as b:
            self.assertEqual(a.namelist(),b.namelist())
            for nombre in a.namelist():
                if nombre!=parte:self.assertEqual(a.read(nombre),b.read(nombre),nombre)
            def sin_formulas(xml):
                for ref in ('K10','R40','R41','R42','U41'):
                    xml=re.sub(rb'(<c\b[^>]*\br="'+ref.encode()+rb'"[^>]*>).*?(</c>)',rb'\1\2',xml,flags=re.S)
                return xml
            self.assertEqual(sin_formulas(a.read(parte)),sin_formulas(b.read(parte)))
        ws=load_workbook(BytesIO(salida),keep_vba=True)['FORMATO']
        self.assertIn('ISNUMBER',ws['R40'].value)
        self.assertIn('COUNT(Y25:Z25)',ws['R41'].value)
        self.assertIn('NO PLÁSTICO',ws['R42'].value)
        self.assertIn('CONCATENATE',ws['U41'].value)
        for flags,esperado in [(('R41','R42'),{'R41':'NO PLÁSTICO','R42':'N.P.'}),(('R40','R41'),{'R40':'NO LÍQUIDO','R41':'NO PLÁSTICO'})]:
            salida=escribir_entradas(original,parte,esperado,celdas_formula_editables=flags)
            salida=reparar_formulas_clasificacion(salida,parte)
            ws=load_workbook(BytesIO(salida),keep_vba=True)['FORMATO']
            for ref,valor in esperado.items():self.assertEqual(ws[ref].value,valor)


if __name__=='__main__':unittest.main()
