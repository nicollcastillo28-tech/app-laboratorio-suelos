"""Ajustes de presentación únicamente en los libros que se descargan."""
from io import BytesIO
import re
import zipfile
from openpyxl.utils.cell import range_boundaries, get_column_letter


def _rango(rect):
    c1, r1, c2, r2 = rect
    a, b = f'{get_column_letter(c1)}{r1}', f'{get_column_letter(c2)}{r2}'
    return a if a == b else f'{a}:{b}'


def _interseccion(a, b):
    c1, r1, c2, r2 = max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3])
    return (c1, r1, c2, r2) if c1 <= c2 and r1 <= r2 else None


def _restar(rect, parte):
    c1, r1, c2, r2 = rect
    x1, y1, x2, y2 = parte
    return [r for r in [(c1,r1,c2,y1-1),(c1,y2+1,c2,r2),(c1,y1,x1-1,y2),(x2+1,y1,c2,y2)]
            if r[0] <= r[2] and r[1] <= r[3]]


def corregir_avisos_vacios(xml):
    """Las reglas de casillas vacías consultan su celda, o el ancla de su combinación.

    Separar los rangos evita la referencia desplazada de las reglas de varias áreas.
    Se mantienen las reglas de resultados, colores, cobertura y orden de precedencia.
    """
    merges = [range_boundaries(r) for r in re.findall(r'<mergeCell\b[^>]*\bref="([^"]+)"', xml)]
    grupos = re.compile(r'<conditionalFormatting\b([^>]*)>(.*?)</conditionalFormatting>', re.S)
    reglas = re.compile(r'<cfRule\b([^>]*?)(?:/>|>(.*?)</cfRule>)', re.S)
    def grupo(m):
        attrs, cuerpo = m.group(1), m.group(2)
        sr = re.search(r'\bsqref="([^"]+)"', attrs)
        if not sr:
            return m.group(0)
        rangos = sr.group(1).split()
        salida = []
        for regla in reglas.finditer(cuerpo):
            ra, contenido = regla.group(1), regla.group(2) or ''
            formula = re.search(r'<formula>(.*?)</formula>', contenido, re.S)
            simple = formula and re.fullmatch(r'LEN\(TRIM\(\$?[A-Z]+\$?\d+\)\)=0', formula.group(1))
            propia = formula and formula.group(1) == 'LEN(TRIM(INDIRECT(ADDRESS(ROW(),COLUMN()))))=0'
            if not (simple or propia) or not re.search(r'\btype="(?:containsBlanks|expression)"', ra):
                salida.append(f'<conditionalFormatting{attrs}>{regla.group(0)}</conditionalFormatting>')
                continue
            destinos = []
            for rango in rangos:
                pendientes = [range_boundaries(rango)]
                for combinado in merges:
                    resto = []
                    for rect in pendientes:
                        parte = _interseccion(rect, combinado)
                        if parte:
                            ancla = f'${get_column_letter(combinado[0])}${combinado[1]}'
                            destinos.append((parte, ancla))
                            resto.extend(_restar(rect, parte))
                        else:
                            resto.append(rect)
                    pendientes = resto
                destinos.extend((rect, f'{get_column_letter(rect[0])}{rect[1]}') for rect in pendientes)
            ra = re.sub(r'\btype="[^"]+"', 'type="expression"', ra)
            for rect, ancla in dict.fromkeys(destinos):
                nuevo = contenido[:formula.start()] + f'<formula>LEN(TRIM({ancla}))=0</formula>' + contenido[formula.end():]
                nuevos_attrs = re.sub(r'\bsqref="[^"]+"', f'sqref="{_rango(rect)}"', attrs)
                salida.append(f'<conditionalFormatting{nuevos_attrs}><cfRule{ra}>{nuevo}</cfRule></conditionalFormatting>')
        return ''.join(salida) if salida else m.group(0)
    resultado = grupos.sub(grupo, xml)
    if resultado == xml:
        return xml
    # Las reglas divididas deben tener prioridades únicas, conservando su orden relativo.
    encontrados = list(re.finditer(r'<cfRule\b[^>]*\bpriority="(\d+)"[^>]*>', resultado))
    orden = {m.start(): i+1 for i,m in enumerate(sorted(encontrados,key=lambda m:(int(m.group(1)),m.start())))}
    return re.sub(r'<cfRule\b[^>]*\bpriority="\d+"[^>]*>',
                  lambda m: re.sub(r'\bpriority="\d+"', f'priority="{orden[m.start()]}"', m.group(0)), resultado)


def ajustar_presentacion(contenido, parte, referencias):
    """Centra las entradas sin modificar valores, formatos numéricos ni objetos."""
    with zipfile.ZipFile(BytesIO(contenido)) as z:
        xml = z.read(parte).decode('utf-8')
        estilos = z.read('xl/styles.xml').decode('utf-8')
        bloque = re.search(r'<cellXfs\b[^>]*>(.*?)</cellXfs>', estilos, re.S)
        xfs = re.findall(r'<xf\b[^>]*?/>|<xf\b[^>]*>.*?</xf>', bloque.group(1), re.S)
        existentes = {xf:i for i,xf in enumerate(xfs)}
        indices = {}
        def centrado(indice):
            if indice in indices:
                return indices[indice]
            xf = xfs[indice]
            tag = re.match(r'<xf\b[^>]*>', xf).group(0)
            attrs = re.sub(r'\s+applyAlignment="[^"]*"', '', tag).rstrip('>/')
            tag_nuevo = attrs + ' applyAlignment="1">'
            cuerpo = '' if tag.endswith('/>') else xf[len(tag):-5]
            al = re.search(r'<alignment\b[^>]*?/>', cuerpo)
            alineacion = al.group(0) if al else '<alignment/>'
            alineacion = re.sub(r'\s+(?:horizontal|vertical|indent|relativeIndent)="[^"]*"', '', alineacion)
            alineacion = alineacion[:-2] + ' horizontal="center" vertical="center"/>'
            if al:
                cuerpo = cuerpo[:al.start()] + alineacion + cuerpo[al.end():]
            else:
                # alignment precede protection y extLst en CT_Xf.
                posicion = re.search(r'<(?:protection|extLst)\b', cuerpo)
                i = posicion.start() if posicion else len(cuerpo)
                cuerpo = cuerpo[:i] + alineacion + cuerpo[i:]
            nuevo = tag_nuevo + cuerpo + '</xf>'
            if nuevo not in existentes:
                existentes[nuevo] = len(xfs)
                xfs.append(nuevo)
            indices[indice] = existentes[nuevo]
            return indices[indice]
        # Identificadores estables: dos muestras del mismo formato comparten estilos
        # al agruparse, aunque una tenga campos vacíos o límites NP/NL.
        for indice in range(len(xfs)):
            centrado(indice)
        refs = set(referencias)
        def celda(m):
            if m.group(1) not in refs:
                return m.group(0)
            tag = m.group(0)
            s = re.search(r'\bs="(\d+)"', tag)
            nuevo = str(centrado(int(s.group(1)) if s else 0))
            if s:
                return re.sub(r'\bs="\d+"', f's="{nuevo}"', tag)
            return tag[:-2] + f' s="{nuevo}"/>' if tag.endswith('/>') else tag[:-1] + f' s="{nuevo}">'
        xml = re.sub(r'<c\b[^>]*\br="([A-Z]+\d+)"[^>]*>', celda, xml)
        xml = corregir_avisos_vacios(xml)
        if len(xfs) != len(re.findall(r'<xf\b', bloque.group(1))):
            apertura = re.sub(r'\bcount="\d+"', f'count="{len(xfs)}"', estilos[bloque.start():bloque.start(1)])
            estilos = estilos[:bloque.start()] + apertura + ''.join(xfs) + '</cellXfs>' + estilos[bloque.end():]
        salida = BytesIO()
        with zipfile.ZipFile(salida, 'w', zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                datos = xml.encode('utf-8') if item.filename == parte else estilos.encode('utf-8') if item.filename == 'xl/styles.xml' else z.read(item.filename)
                out.writestr(item, datos)
        return salida.getvalue()


def reparar_formulas_clasificacion(contenido, parte):
    """Corrige errores de texto en límites y CONCAT sin cambiar el diseño ni el método."""
    formulas = {
        'K10': 'CONCATENATE(L9,"-",M9,"-",IF(D12="SONDEO","S",IF(D12="APIQUE","AP",IF(D12="CANTERA","F",""))),F12,"-M",H12)',
        'R40': 'IF(COUNT(S19:U19)=0,"",IF(T19=25,IF(ISNUMBER(T25),ROUNDDOWN(T25,0),""),IF(COUNT(S25:U25)=0,"",ROUNDDOWN(Q52,0))))',
        'R41': 'IF(COUNT(Y25:Z25)=0,"",ROUNDDOWN(AVERAGE(Y25:Z25),0))',
        'R42': 'IF(R40="NO LÍQUIDO","N.L. - N.P.",IF(R41="NO PLÁSTICO","N.P.",IF(AND(ISNUMBER(R40),ISNUMBER(R41)),R40-R41,"")))',
        'U41': 'IFERROR(CONCATENATE(AE41," (",IF(AE41="A-2-6 O A-2-7",AF43,AE43),")"),"-")',
    }
    import html
    with zipfile.ZipFile(BytesIO(contenido)) as z:
        xml = z.read(parte).decode('utf-8')
        for ref, formula in formulas.items():
            patron = re.compile(r'(<c\b[^>]*\br="'+ref+r'"[^>]*>)(.*?)(</c>)', re.S)
            def cambiar(m):
                if not re.search(r'<f\b', m.group(2)):
                    return m.group(0)  # NP/NL ya digitado tiene precedencia.
                cuerpo = re.sub(r'<f\b[^>]*>.*?</f>', '<f>'+html.escape(formula,quote=False)+'</f>', m.group(2), flags=re.S)
                cuerpo = re.sub(r'<v(?:\s[^>]*)?>.*?</v>|<v\s*/>', '', cuerpo, flags=re.S)
                return m.group(1)+cuerpo+m.group(3)
            xml, n = patron.subn(cambiar, xml)
            if n != 1:
                raise ValueError(f'La celda {ref} del formato de clasificación cambió')
        salida=BytesIO()
        with zipfile.ZipFile(salida,'w',zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                out.writestr(item,xml.encode('utf-8') if item.filename==parte else z.read(item.filename))
        return salida.getvalue()
